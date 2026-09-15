#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
司机轨迹坡度分析服务(FastAPI + 本地 30m DEM)。

启动(在 src/ 目录下执行):
  uvicorn server:app --host 0.0.0.0 --port 8107

环境变量:
  DEM_DIR         DEM 瓦片目录(默认 <项目根>/data/dem)
  AUTO_DOWNLOAD   1 = 遇到缺失瓦片时自动从 AWS 下载(需联网)
"""
import csv
import datetime as dt
import io
import math
import os
import re
import sys
import threading
from typing import Any, Dict, List, Literal, Optional, Tuple

import numpy as np
import rasterio
from fastapi import FastAPI, File, Form, HTTPException, UploadFile
from fastapi.concurrency import run_in_threadpool
from fastapi.responses import FileResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel, Field, field_validator

_SRC_DIR = os.path.dirname(os.path.abspath(__file__))    # src/
PROJECT_DIR = os.path.dirname(_SRC_DIR)                  # 仓库根(开发布局)
# PyInstaller 冻结包里 __file__ 与打包的 static 同在 _internal/,不能上跳一层
IS_FROZEN = getattr(sys, "frozen", False)
DEM_DIR = os.environ.get("DEM_DIR", os.path.join(PROJECT_DIR, "data", "dem"))
STATIC_DIR = os.path.join(_SRC_DIR if IS_FROZEN else PROJECT_DIR, "static")
# 按需下载默认开启(数据源免费无需账号);设 AUTO_DOWNLOAD=0 显式关闭
AUTO_DOWNLOAD = os.environ.get("AUTO_DOWNLOAD", "1") != "0"

TILE_RE = re.compile(r"N(\d{2})_00_E(\d{3})_00_DEM")
# 与 download_dem.py 的瓦片网格一致
TILE_GRID = ((17, 54), (73, 136))  # [lat_min, lat_max), [lon_min, lon_max)

EARTH_R = 6371008.8


# ---------------- GCJ-02 -> WGS-84 ----------------
# 国内 GPS/北斗平台回传的坐标常为 GCJ-02(火星坐标),相对 WGS-84 偏移 50~500m,
# 不转换会导致山区高程采错位置。

_GCJ_A = 6378245.0
_GCJ_EE = 0.00669342162296594323


def _transform_lat(x, y):
    ret = -100.0 + 2.0 * x + 3.0 * y + 0.2 * y * y + 0.1 * x * y + 0.2 * math.sqrt(abs(x))
    ret += (20.0 * math.sin(6.0 * x * math.pi) + 20.0 * math.sin(2.0 * x * math.pi)) * 2.0 / 3.0
    ret += (20.0 * math.sin(y * math.pi) + 40.0 * math.sin(y / 3.0 * math.pi)) * 2.0 / 3.0
    ret += (160.0 * math.sin(y / 12.0 * math.pi) + 320.0 * math.sin(y * math.pi / 30.0)) * 2.0 / 3.0
    return ret


def _transform_lon(x, y):
    ret = 300.0 + x + 2.0 * y + 0.1 * x * x + 0.1 * x * y + 0.1 * math.sqrt(abs(x))
    ret += (20.0 * math.sin(6.0 * x * math.pi) + 20.0 * math.sin(2.0 * x * math.pi)) * 2.0 / 3.0
    ret += (20.0 * math.sin(x * math.pi) + 40.0 * math.sin(x / 3.0 * math.pi)) * 2.0 / 3.0
    ret += (150.0 * math.sin(x / 12.0 * math.pi) + 300.0 * math.sin(x / 30.0 * math.pi)) * 2.0 / 3.0
    return ret


def wgs84_to_gcj02(lng, lat):
    if lng < 72.004 or lng > 137.8347 or lat < 0.8293 or lat > 55.8271:
        return lng, lat
    dlat = _transform_lat(lng - 105.0, lat - 35.0)
    dlng = _transform_lon(lng - 105.0, lat - 35.0)
    radlat = lat / 180.0 * math.pi
    magic = math.sin(radlat)
    magic = 1 - _GCJ_EE * magic * magic
    sqrtmagic = math.sqrt(magic)
    dlat = (dlat * 180.0) / ((_GCJ_A * (1 - _GCJ_EE)) / (magic * sqrtmagic) * math.pi)
    dlng = (dlng * 180.0) / (_GCJ_A / sqrtmagic * math.cos(radlat) * math.pi)
    return lng + dlng, lat + dlat


def gcj02_to_wgs84(lng, lat):
    wlng, wlat = lng, lat
    for _ in range(5):  # 迭代逼近,误差 < 1e-6 度
        glng, glat = wgs84_to_gcj02(wlng, wlat)
        wlng += lng - glng
        wlat += lat - glat
    return wlng, wlat


# ---------------- DEM 瓦片池 ----------------

class DemPool:
    def __init__(self, dem_dir: str):
        self.dem_dir = dem_dir
        self.tiles: Dict[Tuple[int, int], str] = {}      # (lat_floor, lon_floor) -> 文件路径
        self._handles: Dict[Tuple[int, int], Any] = {}
        self._sizes: Dict[Tuple[int, int], int] = {}     # 打开时的文件大小,用于检测运行中被删/被改
        self._lock = threading.Lock()
        self.scan()

    def scan(self):
        if not os.path.isdir(self.dem_dir):
            os.makedirs(self.dem_dir, exist_ok=True)
            return
        for fn in os.listdir(self.dem_dir):
            if not fn.endswith(".tif"):
                continue
            m = TILE_RE.search(fn)
            if m:
                self.tiles[(int(m.group(1)), int(m.group(2)))] = os.path.join(self.dem_dir, fn)

    @staticmethod
    def tile_of(lon: float, lat: float):
        (lat_min, lat_max), (lon_min, lon_max) = TILE_GRID
        if not (lat_min <= lat < lat_max and lon_min <= lon < lon_max):
            return None
        return math.floor(lat), math.floor(lon)

    def missing(self, points) -> List[Tuple[int, int]]:
        need = {t for lon, lat in points if (t := self.tile_of(lon, lat)) is not None}
        return sorted(t for t in need if t not in self.tiles)

    def sample(self, points: List[Tuple[float, float]]) -> np.ndarray:
        """按输入顺序返回每个点的高程(m),无数据处为 nan。"""
        elevs = np.full(len(points), np.nan)
        groups: Dict[Tuple[int, int], List[int]] = {}
        for i, pt in enumerate(points):
            t = self.tile_of(*pt)
            if t is not None:
                groups.setdefault(t, []).append(i)
        with self._lock:
            for t, idxs in groups.items():
                if t not in self._handles:
                    self._handles[t] = rasterio.open(self.tiles[t])   # 打不开则向上抛,触发恢复流程
                    self._sizes[t] = os.path.getsize(self.tiles[t])
                elif os.path.getsize(self.tiles[t]) != self._sizes[t]:
                    raise RuntimeError(f"DEM 瓦片在服务运行期间被修改: {self.tiles[t]}")
                ds = self._handles[t]
                vals = [v[0] for v in ds.sample([points[i] for i in idxs])]
                nodata = ds.nodata
                for i, v in zip(idxs, vals):
                    if v is None or (isinstance(v, float) and math.isnan(v)):
                        continue
                    if nodata is not None and v == nodata:
                        continue
                    elevs[i] = float(v)
        return elevs


POOL = DemPool(DEM_DIR)


_DL_LOCK = threading.Lock()   # _DL_STATE 的短临界区锁(状态读写,绝不能长持)
_DL_ONCE = threading.Lock()   # 串行化补下载本身:并发分析请求只触发一次下载
# 前端经 /download/status 轮询:done/total 为瓦片数,downloading 为正在下载的瓦片
# (label -> [已收字节, 总字节]),bytes_done/bytes_total 为字节级总进度
_DL_STATE = {
    "active": False, "done": 0, "total": 0, "failed": [],
    "downloading": {}, "bytes_done": 0, "bytes_total": 0,
}
_EST_TILE_BYTES = 20 * 1024 * 1024   # HEAD 探测体积失败时的单片估算值


def _content_length(url: str) -> int:
    """HEAD 探测瓦片体积,失败返回 0(调用方按估算值兜底)。"""
    import urllib.request
    try:
        req = urllib.request.Request(url, method="HEAD")
        with urllib.request.urlopen(req, timeout=10) as r:
            return int(r.headers.get("Content-Length") or 0)
    except Exception:
        return 0


def _download_missing(points):
    miss = POOL.missing(points)
    if not miss:
        return
    if not AUTO_DOWNLOAD:
        names = ", ".join(f"Copernicus_DSM_COG_10_N{a:02d}_00_E{b:03d}_00_DEM" for a, b in miss)
        raise HTTPException(
            422,
            detail=f"轨迹覆盖 {len(miss)} 个本地缺失的 DEM 瓦片: {names}。"
                   f"请运行 download_dem.py 下载,或不要设置 AUTO_DOWNLOAD=0(默认开启自动下载)。",
        )
    from concurrent.futures import ThreadPoolExecutor, as_completed
    from download_dem import download_tile, tile_url

    def label_of(a, b):
        return f"N{a:02d}_E{b:03d}"

    def on_chunk(label, received, total):
        with _DL_LOCK:
            entry = _DL_STATE["downloading"].get(label)
            if entry is not None:
                entry[0] = received
                if total > 0:
                    entry[1] = total

    with _DL_ONCE:                       # 并发请求只下载一次
        miss = POOL.missing(points)      # double-check
        if not miss:
            return
        with _DL_LOCK:
            _DL_STATE.update(active=True, done=0, total=len(miss), failed=[],
                             downloading={}, bytes_done=0,
                             bytes_total=len(miss) * _EST_TILE_BYTES)
        try:
            # 先探测各瓦片真实体积,换算出字节级总进度;探测失败沿用估算值
            with ThreadPoolExecutor(max_workers=8) as hx:
                totals = list(hx.map(lambda ab: _content_length(tile_url(*ab)), miss))
            sizes = {(a, b): (n or _EST_TILE_BYTES) for (a, b), n in zip(miss, totals)}
            with _DL_LOCK:
                _DL_STATE["bytes_total"] = sum(sizes.values())

            def job(a, b):
                label = label_of(a, b)
                with _DL_LOCK:
                    _DL_STATE["downloading"][label] = [0, sizes[(a, b)]]
                st = "fail"
                try:
                    st, _, size = download_tile(
                        a, b, DEM_DIR, progress=lambda recv, tot: on_chunk(label, recv, tot))
                finally:
                    with _DL_LOCK:
                        _DL_STATE["downloading"].pop(label, None)
                        _DL_STATE["done"] += 1
                        if st in ("ok", "skip"):
                            _DL_STATE["bytes_done"] += size or sizes[(a, b)]
                return st

            with ThreadPoolExecutor(max_workers=4) as ex:
                futs = {ex.submit(job, a, b): (a, b) for a, b in miss}
                for fut in as_completed(futs):
                    st = fut.result()
                    if st == "fail":
                        with _DL_LOCK:
                            _DL_STATE["failed"].append(label_of(*futs[fut]))
            POOL.scan()
        finally:
            with _DL_LOCK:
                _DL_STATE["active"] = False
        if _DL_STATE["failed"]:
            names = ", ".join(_DL_STATE["failed"])
            raise HTTPException(502, f"自动下载瓦片失败: {names}(网络问题)。请点「重试」继续(已下好的不会重复下载),或手动运行 download_dem.py")


def download_status():
    with _DL_LOCK:
        s = dict(_DL_STATE)
        s["failed"] = list(_DL_STATE["failed"])
        s["downloading"] = {k: list(v) for k, v in sorted(_DL_STATE["downloading"].items())}
    return s


def _purge_broken_tiles(points):
    """删除无法打开的瓦片文件(下载中断或运行中被改动留下的坏文件)。"""
    for t in {POOL.tile_of(*p) for p in points}:
        if t is None:
            continue
        path = POOL.tiles.get(t)
        if not path or not os.path.exists(path):
            with POOL._lock:
                POOL._handles.pop(t, None)
                POOL.tiles.pop(t, None)
            continue
        try:
            rasterio.open(path).close()
        except Exception:
            with POOL._lock:
                POOL._handles.pop(t, None)
                POOL.tiles.pop(t, None)
            try:
                os.remove(path)
            except OSError:
                pass


def sample_elevations(points):
    """确保瓦片就绪并采样;若因坏文件失败,清理后重下一次再采。"""
    _download_missing(points)
    try:
        return POOL.sample(points)
    except Exception:
        _purge_broken_tiles(points)
        _download_missing(points)
        return POOL.sample(points)


# ---------------- 指标计算 ----------------

def haversine_m(lon1, lat1, lon2, lat2):
    lon1, lat1, lon2, lat2 = map(np.radians, (lon1, lat1, lon2, lat2))
    dlat = lat2 - lat1
    dlon = lon2 - lon1
    a = np.sin(dlat / 2) ** 2 + np.cos(lat1) * np.cos(lat2) * np.sin(dlon / 2) ** 2
    return 2 * EARTH_R * np.arcsin(np.sqrt(a))


def nan_moving_average(arr: np.ndarray, k: int) -> np.ndarray:
    if k <= 1:
        return arr
    kernel = np.ones(min(k, len(arr))) / min(k, len(arr))
    mask = (~np.isnan(arr)).astype(float)
    filled = np.where(mask > 0, arr, 0.0)
    s = np.convolve(filled, kernel, mode="same")
    n = np.convolve(mask, kernel, mode="same")
    return np.divide(s, n, out=np.zeros_like(s), where=n > 0)


MAX_GRADE_WINDOW_M = 100.0
MAX_POINTS = 300_000


def _extreme_grades(cum, elevs, gd, valid, d):
    """最大爬坡/下坡坡度:按 >=100m 滑动窗口计算,抑制 DSM 在桥梁/水面/建筑处的单点跳变。
    轨迹总长不足 100m 时退化为逐段(>=5m)统计。"""
    js = np.searchsorted(cum, cum + MAX_GRADE_WINDOW_M, side="left")
    ok = js < len(cum)
    if ok.any():
        ii, jj = np.nonzero(ok)[0], js[ok]
        with np.errstate(invalid="ignore", divide="ignore"):
            g = (elevs[jj] - elevs[ii]) / (cum[jj] - cum[ii])
        g = g[np.isfinite(g)]
    else:
        g = gd[valid & (d >= 5.0)]
    pos, neg = g[g > 0], g[g < 0]
    return (float(pos.max()) if pos.size else 0.0, float(neg.min()) if neg.size else 0.0)


def _haversine_scalar(p0, p1):
    lon1, lat1, lon2, lat2 = map(math.radians, (p0[0], p0[1], p1[0], p1[1]))
    a = math.sin((lat2 - lat1) / 2) ** 2 + math.cos(lat1) * math.cos(lat2) * math.sin((lon2 - lon1) / 2) ** 2
    return 2 * EARTH_R * math.asin(math.sqrt(a))


def densify(orig, pts, step_m: float, cum: Optional[List[float]] = None):
    """相邻点间距超过 step_m 时按直线线性插入中间点(稀疏轨迹加密,让 DEM 采样能反映沿途地形)。
    返回 (orig2, pts2, src, cum2):src[i] 为原始记录下标,插值点为 None;
    cum 为随点单调递增的累计量(如里程表,米),插值点按线性内插。"""
    n = len(pts)
    if step_m <= 0 or n < 2:
        return list(orig), list(pts), list(range(n)), (list(cum) if cum else None)
    o2, p2, src, c2 = [], [], [], []
    for i in range(n - 1):
        o2.append(orig[i]); p2.append(pts[i]); src.append(i)
        if cum:
            c2.append(cum[i])
        k = int(_haversine_scalar(pts[i], pts[i + 1]) // step_m)
        if k <= 0:
            continue
        if len(p2) + k > MAX_POINTS:
            raise HTTPException(400, f"加密后点数超过 {MAX_POINTS},请增大 densify_m 或拆分轨迹")
        for s in range(1, k + 1):
            t = s / (k + 1)
            o2.append((orig[i][0] + (orig[i + 1][0] - orig[i][0]) * t, orig[i][1] + (orig[i + 1][1] - orig[i][1]) * t))
            p2.append((pts[i][0] + (pts[i + 1][0] - pts[i][0]) * t, pts[i][1] + (pts[i + 1][1] - pts[i][1]) * t))
            src.append(None)
            if cum:
                c2.append(cum[i] + (cum[i + 1] - cum[i]) * t)
    o2.append(orig[-1]); p2.append(pts[-1]); src.append(n - 1)
    if cum:
        c2.append(cum[-1])
    return o2, p2, src, c2


def analyze(
    lons: np.ndarray,
    lats: np.ndarray,
    elevs: np.ndarray,
    grade_threshold: float,
    smooth_window_m: float,
    with_detail: bool = False,
    dist_override: Optional[np.ndarray] = None,
):
    if dist_override is not None and len(dist_override) != len(lons) - 1:
        raise HTTPException(400, "dist_override 长度与轨迹点数不匹配")
    if dist_override is not None:
        d = np.asarray(dist_override, dtype=float)   # 权威分段距离(如里程表增量),米
    else:
        d = haversine_m(lons[:-1], lats[:-1], lons[1:], lats[1:])

    # 高程序列平滑:抑制 DSM 植被/建筑噪声与 GPS 定位抖动(窗口按弧长换算点数)
    if smooth_window_m > 0 and len(elevs) >= 5:
        spacing = max(float(np.median(d)), 1.0)
        elevs = nan_moving_average(elevs, int(round(smooth_window_m / spacing)))

    dh = np.diff(elevs)

    valid = d >= 1.0                     # 剔除静止/重复点
    gd = np.zeros_like(d)
    np.divide(dh, d, out=gd, where=valid)  # 坡度 = 高差 / 水平距离

    climb = valid & (gd > grade_threshold)
    descent = valid & (gd < -grade_threshold)
    unknown = valid & np.isnan(gd)          # 高程缺失路段:不计入平路,单独统计
    flat = valid & ~climb & ~descent & ~unknown

    total = float(d[valid].sum())
    climb_d = float(d[climb].sum())
    descent_d = float(d[descent].sum())
    flat_d = float(d[flat].sum())
    unknown_d = float(d[unknown].sum())

    dh_v = dh[valid]
    cum = np.concatenate([[0.0], np.cumsum(np.where(valid, d, 0.0))])
    max_up, min_dn = _extreme_grades(cum, elevs, gd, valid, d)

    # 坡度区间里程分布(%),仅统计高程已知路段
    known = valid & ~np.isnan(gd)
    _edges = [-np.inf, -6, -4, -2, 2, 4, 6, np.inf]
    _hist, _ = np.histogram(gd[known] * 100, bins=_edges, weights=d[known]) if known.any() else ([0.0] * 7, None)

    # 连续长坡(爬坡/下坡段连续合并),按长度取 Top5
    def _runs(mask):
        out, i, n = [], 0, len(mask)
        while i < n:
            if not mask[i]:
                i += 1
                continue
            j = i
            while j < n and mask[j]:
                j += 1
            dist = float(d[i:j].sum())
            if dist >= 200:
                dh_r = float(elevs[j] - elevs[i])
                out.append({"i0": int(i), "i1": int(j), "start_km": round(float(cum[i]) / 1000, 2),
                            "dist_m": round(dist, 1), "dh_m": round(dh_r, 1),
                            "avg_grade": round(dh_r / dist, 4)})
            i = j
        return sorted(out, key=lambda r: -r["dist_m"])[:5]

    result = {
        "point_count": int(len(lons)),
        "total_distance_m": round(total, 1),
        "climb_distance_m": round(climb_d, 1),
        "descent_distance_m": round(descent_d, 1),
        "flat_distance_m": round(flat_d, 1),
        "unknown_distance_m": round(unknown_d, 1),
        "climb_ratio": round(climb_d / total, 4) if total else None,
        "descent_ratio": round(descent_d / total, 4) if total else None,
        "flat_ratio": round(flat_d / total, 4) if total else None,
        "unknown_ratio": round(unknown_d / total, 4) if total else None,
        "elevation_coverage": round(1 - unknown_d / total, 4) if total else None,
        "total_ascent_m": round(float(dh_v[dh_v > 0].sum()), 1),
        "total_descent_m": round(float(-dh_v[dh_v < 0].sum()), 1),
        "ascent_per_100km_m": round(float(dh_v[dh_v > 0].sum()) * 100000 / total, 1) if total else None,
        "max_climb_grade": round(max_up, 4),      # 正值,0.02 = 2%
        "max_descent_grade": round(-min_dn, 4),   # 正值表示下坡幅度
        "min_grade": round(min_dn, 4),
        "elev_min_m": round(float(np.nanmin(elevs)), 1),
        "elev_max_m": round(float(np.nanmax(elevs)), 1),
        "grade_hist": {"bins": ["<-6", "-6~-4", "-4~-2", "±2", "2~4", "4~6", ">6"],
                       "dist_m": [round(float(x), 1) for x in _hist]},
        "top_climbs": _runs(climb),
        "top_descents": _runs(descent),
    }
    if with_detail:
        # 逐点明细:累计里程、平滑后高程、所在段坡度(第 i 点记录 i-1→i 段;静止段记 0;高程缺失为 null)
        pg = np.concatenate([[0.0], gd])
        result["detail"] = {
            "cum_distance_m": [round(float(x), 1) for x in cum],
            "elevations": [None if np.isnan(e) else round(float(e), 1) for e in elevs],
            "grades": [None if not np.isfinite(g) else round(float(g), 4) for g in pg],
        }
    return result


# ---------------- API ----------------

app = FastAPI(title="司机轨迹坡度分析服务", version="1.0.0")


class AnalyzeRequest(BaseModel):
    # 轨迹点:[ [lon, lat], ... ] 或 [ {"lon":..,"lat":..}, ... ],按时间顺序
    points: List[Any]
    coord_system: Literal["wgs84", "gcj02"] = "wgs84"
    grade_threshold: float = Field(0.02, gt=0, le=0.5, description="坡度阈值,0.02 = 2%")
    smooth_window_m: float = Field(150.0, ge=0, description="高程平滑窗口(米),0 关闭")
    densify_m: float = Field(0.0, ge=0, le=1000, description="沿线加密步长(米),0 关闭;点距稀疏时建议 30")

    @field_validator("points")
    @classmethod
    def _parse_points(cls, v):
        pts = []
        for p in v:
            if isinstance(p, dict):
                pts.append((float(p["lon"]), float(p["lat"])))
            elif isinstance(p, (list, tuple)) and len(p) >= 2:
                pts.append((float(p[0]), float(p[1])))
            else:
                raise ValueError("每个点应为 [lon, lat] 或 {lon, lat}")
        if len(pts) < 2:
            raise ValueError("至少需要 2 个轨迹点")
        return pts


@app.post("/analyze")
def analyze_route(req: AnalyzeRequest):
    if req.coord_system == "gcj02":
        points = [gcj02_to_wgs84(lon, lat) for lon, lat in req.points]
    else:
        points = req.points
    _, points, _, _ = densify(points, points, req.densify_m)

    elevs = sample_elevations(points)

    nan_count = int(np.isnan(elevs).sum())
    if nan_count > len(elevs) * 0.5:
        raise HTTPException(422, detail=f"{nan_count}/{len(elevs)} 个点取不到高程(轨迹可能超出数据覆盖范围)")

    lons = np.array([p[0] for p in points])
    lats = np.array([p[1] for p in points])
    result = analyze(lons, lats, elevs, req.grade_threshold, req.smooth_window_m)
    result["nan_point_count"] = nan_count
    return result


@app.get("/download/status")
def dl_status():
    """DEM 补下载进度(前端在分析请求期间轮询展示)。"""
    return download_status()


@app.get("/health")
def health():
    lats = [a for a, _ in POOL.tiles]
    lons = [b for _, b in POOL.tiles]
    return {
        "tile_count": len(POOL.tiles),
        "dem_dir": DEM_DIR,
        "auto_download": AUTO_DOWNLOAD,
        "coverage": {
            "lat": [min(lats), max(lats) + 1] if lats else None,
            "lon": [min(lons), max(lons) + 1] if lons else None,
        },
    }


# ---------------- 文件上传(xlsx / csv) ----------------

LON_KEYS = ("经度", "lon", "lng", "longitude")
LAT_KEYS = ("纬度", "lat", "latitude")
TIME_KEYS = ("记录时间", "时间", "time", "datetime", "timestamp")
MILE_KEYS = ("里程", "mileage", "odometer", "distance")


def _find_col(headers: List[Any], keys) -> Optional[int]:
    hs = [str(h).strip().lower() if h is not None else "" for h in headers]
    for k in keys:
        if k in hs:
            return hs.index(k)
    for k in keys:
        for i, h in enumerate(hs):
            if k in h:
                return i
    return None


def _parse_time(v) -> Optional[dt.datetime]:
    if v is None or v == "":
        return None
    if isinstance(v, dt.datetime):
        return v
    if isinstance(v, (int, float)):
        # 兼容秒 / 毫秒时间戳
        return dt.datetime.fromtimestamp(v / 1000 if v > 1e11 else v)
    s = str(v).strip().replace("/", "-").replace("T", " ")
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d %H:%M:%S.%f"):
        try:
            return dt.datetime.strptime(s, fmt)
        except ValueError:
            continue
    return None


def _to_float(v) -> Optional[float]:
    try:
        f = float(v)
        return None if math.isnan(f) else f
    except (TypeError, ValueError):
        return None


def _read_table(filename: str, content: bytes) -> List[List[Any]]:
    name = filename.lower()
    if name.endswith((".xlsx", ".xlsm", ".xls")):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        return [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
    for enc in ("utf-8-sig", "gbk"):
        try:
            text = content.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise HTTPException(400, "CSV 编码无法识别,请使用 UTF-8 或 GBK")
    dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t") if text.strip() else csv.excel
    return [row for row in csv.reader(io.StringIO(text), dialect)]


def parse_upload(filename: str, content: bytes) -> List[Dict[str, Any]]:
    """解析上传文件,返回按时间升序排列的记录:{lon, lat, time}(坐标为文件原始坐标系)。"""
    rows = _read_table(filename, content)
    rows = [r for r in rows if r and any(c not in (None, "") for c in r)]
    if len(rows) < 3:
        raise HTTPException(400, "文件内容为空或行数不足")

    headers = rows[0]
    ilon, ilat = _find_col(headers, LON_KEYS), _find_col(headers, LAT_KEYS)
    if ilon is None or ilat is None:
        raise HTTPException(
            400, f"未识别到经纬度列。表头应包含「经度/纬度」或 lon/lng/lat,当前表头: {headers}"
        )
    itime = _find_col(headers, TIME_KEYS)
    imile = _find_col(headers, MILE_KEYS)

    records = []
    for idx, r in enumerate(rows[1:]):
        lon = _to_float(r[ilon]) if ilon < len(r) else None
        lat = _to_float(r[ilat]) if ilat < len(r) else None
        if lon is None or lat is None:
            continue
        if not (-180 <= lon <= 180 and -90 <= lat <= 90):
            continue
        t = _parse_time(r[itime]) if itime is not None and itime < len(r) else None
        mile = _to_float(r[imile]) if imile is not None and imile < len(r) else None
        records.append({"lon": lon, "lat": lat, "time": t, "mile": mile, "_idx": idx})

    if len(records) < 2:
        raise HTTPException(400, "有效轨迹点不足 2 个")

    # 有时间列则按时间升序(文件常为倒序导出);无时间的行保持相对顺序排在末尾
    if itime is not None and sum(1 for r in records if r["time"]) >= 2:
        records.sort(key=lambda r: (r["time"] is None, r["time"] or dt.datetime.min, r["_idx"]))
    return records


@app.post("/analyze/file")
async def analyze_file(
    file: UploadFile = File(...),
    coord_system: str = Form("gcj02"),
    grade_threshold: float = Form(0.02),
    smooth_window_m: float = Form(150.0),
    densify_m: float = Form(0.0),
    distance_source: str = Form("auto"),
):
    if coord_system not in ("wgs84", "gcj02"):
        raise HTTPException(400, "coord_system 只能为 wgs84 或 gcj02")
    if distance_source not in ("auto", "odometer", "gps"):
        raise HTTPException(400, "distance_source 只能为 auto / odometer / gps")
    if not (0 < grade_threshold <= 0.5) or smooth_window_m < 0 or not (0 <= densify_m <= 1000):
        raise HTTPException(400, "参数超出范围")

    filename = file.filename or ""
    content = await file.read()
    # 解析 + 可能的 DEM 补下载 + 采样都在线程池执行,不阻塞事件循环;
    # 否则下载期间 /download/status 无响应,前端进度条不会动
    return await run_in_threadpool(
        _analyze_file_impl, filename, content,
        coord_system, grade_threshold, smooth_window_m, densify_m, distance_source)


def _analyze_file_impl(
    filename: str, content: bytes, coord_system: str,
    grade_threshold: float, smooth_window_m: float, densify_m: float, distance_source: str,
):
    records = parse_upload(filename, content)
    orig = [(r["lon"], r["lat"]) for r in records]
    points = [gcj02_to_wgs84(*p) for p in orig] if coord_system == "gcj02" else orig

    # 原始点距(排除静止重复点):远大于 DEM 分辨率(30m)时坡度指标不可靠,前端据此提示
    _p = np.array(points)
    _d = haversine_m(_p[:-1, 0], _p[:-1, 1], _p[1:, 0], _p[1:, 1])
    _dv = _d[_d >= 1.0]
    mean_spacing = round(float(_dv.mean()), 1) if _dv.size else 0.0
    gps_chord_m = round(float(_dv.sum()), 1)

    # 时间统计与轨迹断点(基于原始记录):总时长、停留、≥30 分钟的记录中断
    times = [r["time"] for r in records]
    vt = [t for t in times if t]
    time_stats = {}
    if len(vt) >= 2:
        time_stats["time_start"] = vt[0].strftime("%Y-%m-%d %H:%M:%S")
        time_stats["time_end"] = vt[-1].strftime("%Y-%m-%d %H:%M:%S")
        time_stats["duration_h"] = round((vt[-1] - vt[0]).total_seconds() / 3600, 1)
        still_min, gaps = 0.0, []
        for k in range(1, len(times)):
            t0, t1 = times[k - 1], times[k]
            if not (t0 and t1):
                continue
            mins = (t1 - t0).total_seconds() / 60
            if mins >= 30:
                gaps.append({"after_time": t0.strftime("%m-%d %H:%M"), "minutes": round(mins)})
            elif _d[k - 1] < 1.0:
                still_min += mins
        time_stats["still_min"] = round(still_min)
        time_stats["time_gaps"] = gaps[:20]

    # 里程来源:文件自带累计里程列(车辆里程表,km)比 GPS 折线更接近真实行驶距离
    # auto = 列有效则用;odometer = 强制,无效报错;gps = 不用
    miles = np.array([r["mile"] if r["mile"] is not None else np.nan for r in records])
    cum_odo = None
    if not np.all(np.isnan(miles)) and distance_source in ("auto", "odometer"):
        dm = np.diff(miles) * 1000.0
        total_odo = float(np.clip(dm[~np.isnan(dm)], 0, None).sum())
        bad = float(np.mean(dm[~np.isnan(dm)] < -5)) if len(dm) else 1.0
        valid_odo = (bad <= 0.05 and total_odo > 100 and
                     0.3 * gps_chord_m <= total_odo <= 3 * gps_chord_m)
        if valid_odo:
            cum_odo = np.maximum.accumulate(
                np.concatenate([[0.0], np.cumsum(np.nan_to_num(dm, nan=0.0))])
            ).tolist()
        elif distance_source == "odometer":
            raise HTTPException(400, "里程列无效(非累计值或与轨迹不符),无法作为距离来源")

    orig, points, src, cum_odo_d = densify(orig, points, densify_m, cum_odo)
    dist_source = "odometer" if cum_odo_d is not None else "gps"
    dist_override = np.diff(cum_odo_d) if cum_odo_d is not None else None

    elevs = sample_elevations(points)
    nan_count = int(np.isnan(elevs).sum())
    if nan_count > len(elevs) * 0.5:
        raise HTTPException(422, f"{nan_count}/{len(elevs)} 个点取不到高程(轨迹可能超出数据覆盖范围)")

    lons = np.array([p[0] for p in points])
    lats = np.array([p[1] for p in points])
    result = analyze(lons, lats, elevs, grade_threshold, smooth_window_m,
                     with_detail=True, dist_override=dist_override)
    detail = result.pop("detail")
    result["nan_point_count"] = nan_count
    result["original_point_count"] = len(records)
    result["mean_spacing_m"] = mean_spacing
    result["distance_source"] = dist_source
    result["gps_chord_m"] = gps_chord_m
    result.update(time_stats)

    rows = []
    for i, s in enumerate(src):
        rec = records[s] if s is not None else None
        rows.append({
            "seq": s + 1 if rec else None,
            "time": rec["time"].strftime("%Y-%m-%d %H:%M:%S") if rec and rec["time"] else None,
            "lon": orig[i][0],               # 原始坐标系(供地图直接展示)
            "lat": orig[i][1],
            "elev": detail["elevations"][i],
            "grade": detail["grades"][i],
            "dist_m": detail["cum_distance_m"][i],
            "interp": rec is None,           # True = 沿线加密插入的点
        })
    return {
        "filename": filename,
        "coord_system": coord_system,
        "grade_threshold": grade_threshold,
        "smooth_window_m": smooth_window_m,
        "densify_m": densify_m,
        "metrics": result,
        "rows": rows,
    }


# ---------------- 前端页面 ----------------

app.mount("/static", StaticFiles(directory=STATIC_DIR), name="static")


@app.get("/")
def index():
    return FileResponse(os.path.join(STATIC_DIR, "index.html"),
                        headers={"Cache-Control": "no-cache"})   # 入口页禁用启发式缓存,避免更新界面后浏览器用旧版
